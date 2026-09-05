"""Untrusted OOXML stays bounded, textual codes and lexical Decimal stay exact."""

import io
import zipfile
from decimal import Decimal
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook

from forge_erp.core.errors import Problem
from forge_erp.modules.catalog_import.application.preview import parse_value
from forge_erp.modules.catalog_import.domain.columns import COLUMNS
from forge_erp.modules.catalog_import.infrastructure import xlsx as x


def rewritten(data, changes=None, extra=None):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    for name, transform in (changes or {}).items():
        parts[name] = transform(parts[name])
    parts.update(extra or {})
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content)
    return target.getvalue()


def sheet_cells(data, **changes):
    def modify(content):
        root = ET.fromstring(content)
        for cell in root.iter(x.NS + "c"):
            if cell.get("r") in changes:
                kind, value = changes[cell.get("r")]
                cell.clear()
                cell.set("r", next(k for k, v in changes.items() if v == (kind, value)))
                cell.set("t", kind)
                ET.SubElement(cell, x.NS + "v").text = value
        return ET.tostring(root)

    return rewritten(data, {"xl/worksheets/sheet1.xml": modify})


def simple():
    return x.workbook_bytes(["编码", "名称"], [["000001", "M6*20"]])


def rejected(data, code, *, filename="资料.xlsx", worksheet=None):
    with pytest.raises(Problem) as err:
        x.parse(data, filename, worksheet)
    assert err.value.code == code, err.value


def test_lexical_numeric_cell_never_passes_float_and_text_identifier_keeps_zeroes():
    data = x.workbook_bytes(["商品编码", "价格"], [["00001234", "0"]])
    data = sheet_cells(data, B2=("n", "12345678901234.000001"))
    sheet = x.parse(data, "../../资料.xlsx")
    cells = sheet.rows[0][1]
    assert sheet.filename == "资料.xlsx"
    assert cells["商品编码"].value == "00001234"
    assert cells["价格"] == x.Cell("12345678901234.000001", "number")
    assert parse_value(cells["价格"], COLUMNS["product-prices"]["价格"]) == Decimal(
        "12345678901234.000001"
    )


@pytest.mark.parametrize(
    "value,kind", [("12345", "number"), ("1", "boolean"), ("2026-09-06", "date")]
)
def test_numeric_boolean_and_date_identifiers_are_never_guessed(value, kind):
    with pytest.raises(Problem) as err:
        parse_value(x.Cell(value, kind), COLUMNS["products"]["条码"])
    assert err.value.code == "UNSAFE_IDENTIFIER_CELL"


@pytest.mark.parametrize(
    "value,kind",
    [
        ("NaN", "text"),
        ("Infinity", "number"),
        ("-Infinity", "number"),
        ("1", "boolean"),
        ("2026-09-06", "date"),
        ("#DIV/0", "error"),
    ],
)
def test_decimal_rejects_nonfinite_and_non_numeric_cell_types(value, kind):
    with pytest.raises(Problem) as err:
        parse_value(x.Cell(value, kind), COLUMNS["product-prices"]["价格"])
    assert err.value.code == "INVALID_DECIMAL"


def test_workbook_formula_is_rejected_even_on_unselected_uppercase_xml_sheet():
    data = rewritten(
        simple(),
        extra={
            "xl/worksheets/hidden.XML": b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c><f>WEBSERVICE("https://invalid.test")</f><v>1</v></c></worksheet>'
        },
    )
    rejected(data, "FORMULA_NOT_ALLOWED")


@pytest.mark.parametrize(
    "extra,code",
    [
        ({"xl/vbaProject.bin": b"macro"}, "UNSAFE_WORKBOOK"),
        ({"xl/embeddings/oleObject1.bin": b"object"}, "UNSAFE_WORKBOOK"),
        ({"xl/externalLinks/externalLink1.xml": b"<externalLink/>"}, "UNSAFE_WORKBOOK"),
        (
            {
                "xl/_rels/extra.RELS": b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>'
            },
            "EXTERNAL_LINK_NOT_ALLOWED",
        ),
        ({"xl/unsafe.XML": b'<!DOCTYPE data [<!ENTITY a "x">]><data>&a;</data>'}, "UNSAFE_XML"),
        ({"xl/unsafe.xml": b"<x>" * 70 + b"t" + b"</x>" * 70}, "UNSAFE_XML"),
        ({"../escape.xml": b"<x/>"}, "UNSAFE_ZIP_PATH"),
        ({"/absolute.xml": b"<x/>"}, "UNSAFE_ZIP_PATH"),
        ({"xl/./escape.xml": b"<x/>"}, "UNSAFE_ZIP_PATH"),
        ({"xl\\escape.xml": b"<x/>"}, "UNSAFE_ZIP_PATH"),
    ],
)
def test_reject_unsafe_ooxml_members_everywhere(extra, code):
    rejected(rewritten(simple(), extra=extra), code)


def test_shared_string_negative_index_is_invalid_and_rich_phonetics_are_excluded():
    data = rewritten(
        sheet_cells(simple(), A2=("s", "-1")),
        extra={
            "xl/sharedStrings.xml": b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>LAST</t></si></sst>'
        },
    )
    rejected(data, "INVALID_XLSX")
    data = rewritten(
        sheet_cells(simple(), A2=("s", "0")),
        extra={
            "xl/sharedStrings.xml": '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>00</t></r><rPh><t>拼音不得混入</t></rPh><r><t>12</t></r></si></sst>'.encode()
        },
    )
    assert x.parse(data, "x.xlsx").rows[0][1]["编码"].value == "0012"


def test_multiple_sheets_require_explicit_selection_and_formula_in_other_sheet_rejected():
    book = Workbook()
    first = book.active
    first.title = "一"
    first.append(["编码", "名称"])
    first.append(["ONE", "一"])
    second = book.create_sheet("二")
    second.append(["编码", "名称"])
    second.append(["TWO", "二"])
    stream = io.BytesIO()
    book.save(stream)
    rejected(stream.getvalue(), "WORKSHEET_REQUIRED")
    assert x.parse(stream.getvalue(), "x.xlsx", "二").rows[0][1]["编码"].value == "TWO"
    rejected(stream.getvalue(), "UNKNOWN_WORKSHEET", worksheet="不存在")
    first["A2"] = "=1+1"
    stream = io.BytesIO()
    book.save(stream)
    rejected(stream.getvalue(), "FORMULA_NOT_ALLOWED", worksheet="二")
    book.close()


def test_ten_thousand_rows_accepted_and_one_more_rejected():
    rows = [["K" + str(i).zfill(6), "名称" + str(i)] for i in range(10000)]
    assert len(x.parse(x.workbook_bytes(["编码", "名称"], rows), "x.xlsx").rows) == 10000
    rejected(x.workbook_bytes(["编码", "名称"], rows + [["EXTRA", "超过行数"]]), "ROW_LIMIT")


def test_compression_actual_member_total_upload_and_column_limits(monkeypatch):
    rejected(rewritten(simple(), extra={"bomb.bin": b"x" * 100000}), "ZIP_LIMIT")
    monkeypatch.setattr(x, "MAX_MEMBER", 100)
    rejected(simple(), "ZIP_LIMIT")
    monkeypatch.setattr(x, "MAX_MEMBER", 50000000)
    monkeypatch.setattr(x, "MAX_EXPANDED", 100)
    rejected(simple(), "ZIP_LIMIT")
    monkeypatch.setattr(x, "MAX_EXPANDED", 100000000)
    monkeypatch.setattr(x, "MAX_UPLOAD", 100)
    rejected(simple(), "FILE_TOO_LARGE")
    monkeypatch.setattr(x, "MAX_UPLOAD", 10000000)
    rejected(x.workbook_bytes([str(i) for i in range(65)], [["v"] * 65]), "COLUMN_LIMIT")


def test_zip_duplicates_count_and_header_or_extension_validation():
    stream = io.BytesIO(simple())
    with zipfile.ZipFile(stream, "a") as package:
        with pytest.warns(UserWarning, match="Duplicate name"):
            package.writestr("xl/workbook.xml", "<x/>")
    rejected(stream.getvalue(), "UNSAFE_ZIP_PATH")
    rejected(rewritten(simple(), extra={f"extra/{i}": b"" for i in range(1025)}), "ZIP_LIMIT")
    rejected(x.workbook_bytes(["编码", "编码"], [["x", "x"]]), "INVALID_HEADERS")
    rejected(x.workbook_bytes(["编码", "名称"], []), "EMPTY_WORKSHEET")
    rejected(simple(), "INVALID_XLSX", filename="macro.xlsm")
    rejected(b"not a zip", "INVALID_XLSX")


def test_safe_download_never_emits_formula_cells():
    data = x.workbook_bytes(["名称"], [["=SUM(A1:A2)"], ["+1"], ["@x"], ["-3"]])
    book = load_workbook(io.BytesIO(data), data_only=False)
    assert [row[0].value for row in list(book.active)[1:]] == ["=SUM(A1:A2)", "+1", "@x", "-3"]
    assert all(row[0].data_type == "s" for row in book.active)
    book.close()


@pytest.mark.parametrize("format", ["00000000", "0.00E+00", "yyyy-mm-dd"])
def test_display_format_does_not_reconstruct_identifier(format):
    book = Workbook()
    sheet = book.active
    sheet.append(["编码", "名称"])
    sheet.append([1234, "数字编码"])
    sheet["A2"].number_format = format
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    parsed = x.parse(stream.getvalue(), "numbers.xlsx")
    with pytest.raises(Problem) as err:
        parse_value(parsed.rows[0][1]["编码"], COLUMNS["brands"]["编码"])
    assert err.value.code == "UNSAFE_IDENTIFIER_CELL"


def test_encrypted_zip_flags_and_cached_formula_rejected():
    data = bytearray(simple())
    # ZIP encryption bit is meaningful in both the local and central directory headers.
    for marker, offset in [(b"PK\x03\x04", 6), (b"PK\x01\x02", 8)]:
        index = data.find(marker)
        data[index + offset] |= 1
    rejected(bytes(data), "ENCRYPTED_WORKBOOK")
    cached = rewritten(
        simple(),
        {
            "xl/worksheets/sheet1.xml": lambda b: b.replace(
                b"</worksheet>", b'<c r="C2"><f>1+1</f><v>2</v></c></worksheet>'
            )
        },
    )
    rejected(cached, "FORMULA_NOT_ALLOWED")


def test_fake_dimension_cannot_hide_actual_rows(monkeypatch):
    data = x.workbook_bytes(["编码", "名称"], [["A", "一"], ["B", "二"]])
    data = rewritten(data, {"xl/worksheets/sheet1.xml": lambda b: b.replace(b"A1:B3", b"A1:A1")})
    monkeypatch.setattr(x, "MAX_ROWS", 1)
    rejected(data, "ROW_LIMIT")


def test_extreme_integer_exponent_and_deep_attribute_json_fail_cleanly():
    with pytest.raises(Problem) as err:
        parse_value(x.Cell("1e1000000", "number"), COLUMNS["supplier-products"]["提前期天数"])
    assert err.value.code == "INVALID_DECIMAL"
    with pytest.raises(Problem) as err:
        parse_value(
            x.Cell("[" * 1500 + "0" + "]" * 1500, "text"), COLUMNS["categories"]["属性模板"]
        )
    assert err.value.code == "INVALID_ATTRIBUTE"
