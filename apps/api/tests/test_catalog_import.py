"""Real API and ordinary Catalog commands prove all ten import resources."""

import io
from decimal import Decimal as D
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from test_catalog import catalog_client as catalog_client
from test_catalog import create

from forge_erp.core.db import sessions
from forge_erp.modules.catalog_import.application import runner
from forge_erp.modules.catalog_import.domain.columns import COLUMNS
from forge_erp.modules.catalog_import.infrastructure.xlsx import workbook_bytes

BASE = "/api/v1/catalog-imports"


async def upload(c, resource, headers, rows, mode="CREATE_ONLY", key=None, data=None):
    return await c.post(
        BASE + "/previews",
        files={"file": ("资料.xlsx", data if data is not None else workbook_bytes(headers, rows))},
        data={"resource": resource, "mode": mode},
        headers={"Idempotency-Key": key or uuid4().hex},
    )


async def details(c, id):
    response = await c.get(BASE + "/" + str(id))
    assert response.status_code == 200, response.text
    return response.json()


async def preview_rows(c, id):
    response = await c.get(BASE + "/" + str(id) + "/rows?page_size=100")
    assert response.status_code == 200, response.text
    return response.json()["items"]


async def confirm(c, id, key=None):
    batch = await details(c, id)
    return await create(
        c,
        "catalog-imports/" + str(id) + "/confirm",
        {"expected_version": batch["version"], "preview_hash": batch["preview_hash"]},
        key,
    )


@pytest.fixture
async def import_refs(catalog_client):
    c = catalog_client
    result = {}
    for resource, code in [
        ("categories", "CAT"),
        ("brands", "BRAND"),
        ("units", "PCS"),
        ("units", "BOX"),
        ("customers", "CUSTOMER"),
        ("suppliers", "SUPPLIER"),
    ]:
        response = await create(c, resource, {"code": code, "name": code})
        assert response.status_code == 201, response.text
        result[code] = response.json()
    product = await create(
        c,
        "products",
        {
            "sku": "ITEM",
            "name": "既有商品",
            "category_id": result["CAT"]["id"],
            "base_unit_id": result["PCS"]["id"],
        },
    )
    assert product.status_code == 201, product.text
    result["ITEM"] = product.json()
    return result


RESOURCE_DATA = {
    "categories": {"编码": "NEW_CAT", "名称": "紧固件", "父分类编码": "CAT"},
    "brands": {"编码": "NEW_BRAND", "名称": "品牌"},
    "units": {"编码": "NEW_UNIT", "名称": "件"},
    "customers": {"编码": "NEW_CUSTOMER", "名称": "客户", "电话": "0010002345"},
    "suppliers": {"编码": "NEW_SUPPLIER", "名称": "供应商", "邮箱": "s*2@example.test"},
    "warehouses": {"编码": "NEW_WAREHOUSE", "名称": "仓库", "地址": "主店"},
    "products": {
        "商品编码": "NEW_ITEM",
        "商品名称": " M6*20 螺栓 ",
        "分类编码": "CAT",
        "基础单位编码": "PCS",
        "条码": "000012340",
        "首选供应商编码": "SUPPLIER",
    },
    "product-units": {"商品编码": "ITEM", "单位编码": "BOX", "换算率": "12.000001"},
    "product-prices": {
        "商品编码": "ITEM",
        "价格类型": "customer",
        "客户编码": "CUSTOMER",
        "价格": "12345678901234.000001",
    },
    "supplier-products": {
        "商品编码": "ITEM",
        "供应商编码": "SUPPLIER",
        "采购单位编码": "PCS",
        "供应商货号": "S00001",
    },
}


@pytest.mark.parametrize("resource", RESOURCE_DATA)
async def test_all_ten_resources_use_catalog_create_and_complete_update(
    catalog_client,
    identities,
    import_refs,
    resource,
):
    c = catalog_client
    data = RESOURCE_DATA[resource]
    headers = list(COLUMNS[resource])
    # The actual downloaded full template can leave optional cells empty on create.
    template = await c.get(BASE + "/templates/" + resource)
    assert template.status_code == 200
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats")
    book = load_workbook(io.BytesIO(template.content))
    assert len(book.sheetnames) == 1
    assert all(cell.comment and "更新" in cell.comment.text for cell in book.active[1])
    book.close()
    response = await upload(c, resource, headers, [[data.get(k, "") for k in headers]])
    assert response.status_code == 201, response.text
    batch = response.json()
    rows = await preview_rows(c, batch["id"])
    assert batch["status"] == "PREVIEW_READY", rows
    assert rows[0]["action"] == "CREATE" and rows[0]["target_id"] is None
    assert (await confirm(c, batch["id"])).status_code == 200
    result = await runner.drain(sessions, identities[0]["org"])
    assert result == {"succeeded": 1, "failed": 0, "purged": 0}, await preview_rows(c, batch["id"])
    completed = await details(c, batch["id"])
    assert completed["status"] == "COMPLETED" and completed["succeeded"] == 1
    row = (await preview_rows(c, batch["id"]))[0]
    record = await c.get("/api/v1/" + resource + "/" + row["target_id"])
    assert record.status_code == 200, record.text
    record = record.json()
    assert row["target_version"] == record["version"] == 1
    if resource == "products":
        assert record["name"] == "M6×20 螺栓" and record["barcode"] == "000012340"
        assert D(record["min_stock_qty"]) == 0
        units = await c.get("/api/v1/product-units", params={"product_id": row["target_id"]})
        assert units.json()["total"] == 1
    if resource == "customers":
        assert record["phone"] == "0010002345" and record["price_tier"] == "standard"
    if resource == "suppliers":
        assert record["email"] == "s*2@example.test"
    if resource == "product-prices":
        assert D(record["price"]) == D(data["价格"])
    # Update only a note, with natural key columns; all unmapped fields must survive.
    from forge_erp.modules.catalog_import.domain.columns import KEYS

    keys = [h for h, item in COLUMNS[resource].items() if item.field in KEYS[resource]]
    changed = await upload(
        c,
        resource,
        keys + ["备注"],
        [[data.get(h, "") for h in keys] + ["只改备注"]],
        mode="UPDATE_EXISTING",
    )
    assert changed.status_code == 201, changed.text
    next_rows = await preview_rows(c, changed.json()["id"])
    assert changed.json()["status"] == "PREVIEW_READY", next_rows
    assert next_rows[0]["expected_target_id"] == row["target_id"]
    assert next_rows[0]["expected_version"] == 1
    assert len(next_rows[0]["cleaned_values"]) >= len(headers)
    assert (await confirm(c, changed.json()["id"])).status_code == 200
    assert (await runner.drain(sessions, identities[0]["org"]))["succeeded"] == 1
    after = (await c.get("/api/v1/" + resource + "/" + row["target_id"])).json()
    assert after["version"] == 2 and after["notes"] == "只改备注"
    for key, value in record.items():
        if key not in {"version", "updated_at", "notes", "search_text"}:
            assert after[key] == value, key


async def test_preview_and_confirmation_http_idempotency_paging_and_safe_download(
    catalog_client,
    identities,
):
    c = catalog_client
    data = workbook_bytes(["编码", "名称"], [["0001", "=1+1"], ["0002", "+SUM(1,2)"]])
    key = uuid4().hex
    first = await upload(c, "brands", [], [], data=data, key=key)
    assert first.status_code == 201, first.text
    assert (await upload(c, "brands", [], [], data=data, key=key)).json() == first.json()
    conflict = await upload(c, "brands", ["编码", "名称"], [["OTHER", "其他"]], key=key)
    assert conflict.status_code == 409 and conflict.json()["code"] == "IDEMPOTENCY_KEY_REUSED"
    id = first.json()["id"]
    before = await c.get(BASE + "/" + id + "/rows?page=9&page_size=1")
    assert before.json()["items"] == [] and before.json()["total"] == 2
    page = await c.get(BASE + "?page=2&page_size=1")
    assert page.json()["items"] == [] and page.json()["total"] == 1
    download = await c.get(BASE + "/" + id + "/result.xlsx")
    book = load_workbook(io.BytesIO(download.content), data_only=False)
    assert all(cell.data_type != "f" for row in book.active for cell in row)
    assert "=1+1" in [cell.value for row in book.active for cell in row]
    book.close()
    batch = await details(c, id)
    body = {"expected_version": 1, "preview_hash": batch["preview_hash"]}
    key = uuid4().hex
    confirmed = await create(c, "catalog-imports/" + id + "/confirm", body, key)
    assert confirmed.status_code == 200, confirmed.text
    assert (
        await create(c, "catalog-imports/" + id + "/confirm", body, key)
    ).json() == confirmed.json()
    conflict = await create(
        c, "catalog-imports/" + id + "/confirm", body | {"expected_version": 2}, key
    )
    assert conflict.status_code == 409
    assert (await runner.drain(sessions, identities[0]["org"]))["succeeded"] == 2
    assert (await runner.drain(sessions, identities[0]["org"]))["succeeded"] == 0
    failed = await c.get(BASE + "/" + id + "/result.xlsx?failed_only=true")
    assert failed.status_code == 409


@pytest.mark.parametrize(
    "resource,headers,values,code",
    [
        ("brands", ["编码", "名称"], [["DUP", "正常"], ["DUP", ""]], "DUPLICATE_IN_FILE"),
        (
            "products",
            ["商品编码", "商品名称", "分类编码", "基础单位编码"],
            [["X", "商品", "OTHER_TENANT", "PCS"]],
            "UNKNOWN_REFERENCE",
        ),
        (
            "product-prices",
            ["商品编码", "价格类型", "价格"],
            [["ITEM", "standard", "0.0000001"]],
            "INVALID_DECIMAL",
        ),
        (
            "product-units",
            ["商品编码", "单位编码", "换算率"],
            [["ITEM", "BOX", "0"]],
            "INVALID_VALUE",
        ),
        ("brands", ["编码", "名称"], [["BRAND", "不覆盖既有品牌"]], "EXISTING_CODE"),
    ],
)
async def test_invalid_preview_is_nonmutating_and_cannot_confirm(
    catalog_client,
    import_refs,
    resource,
    headers,
    values,
    code,
):
    c = catalog_client
    response = await upload(c, resource, headers, values)
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "PREVIEW_INVALID"
    rows = await preview_rows(c, response.json()["id"])
    assert code in {e["code"] for row in rows for e in row["errors"]}, rows
    blocked = await confirm(c, response.json()["id"])
    assert blocked.status_code == 409 and blocked.json()["code"] == "PREVIEW_HAS_ERRORS"
    failed = await c.get(BASE + "/" + response.json()["id"] + "/result.xlsx?failed_only=true")
    book = load_workbook(io.BytesIO(failed.content))
    assert [cell.value for cell in book.active[1]] == headers
    assert book.active.max_row == len(rows) + 1
    book.close()


@pytest.mark.parametrize(
    "header", ["organization_id", "id", "active", "version", "updated_at", "名称2"]
)
async def test_unknown_and_system_columns_rejected(catalog_client, header):
    result = await upload(
        catalog_client, "brands", ["编码", "名称", header], [["X", "商品", "bad"]]
    )
    assert result.status_code == 422 and result.json()["code"] == "UNKNOWN_COLUMN"


async def test_frozen_update_and_reference_versions_require_new_preview(
    catalog_client, identities, import_refs
):
    c = catalog_client
    response = await upload(
        c, "brands", ["编码", "名称"], [["BRAND", "导入的新名称"]], mode="UPDATE_EXISTING"
    )
    assert response.json()["status"] == "PREVIEW_READY"
    original = import_refs["BRAND"]
    manual = await c.put(
        "/api/v1/brands/" + original["id"],
        json={"code": "BRAND", "name": "人工优先修改", "expected_version": 1},
        headers={"Idempotency-Key": uuid4().hex},
    )
    assert manual.status_code == 200
    assert (await confirm(c, response.json()["id"])).status_code == 200
    assert (await runner.drain(sessions, identities[0]["org"]))["failed"] == 1
    row = (await preview_rows(c, response.json()["id"]))[0]
    assert row["error_code"] == "VERSION_CONFLICT" and not row["retryable"]
    retry = await create(
        c,
        "catalog-imports/" + response.json()["id"] + "/retry",
        {"expected_version": 2, "row_ids": [row["id"]]},
    )
    assert retry.status_code == 409
    new_product = await upload(
        c,
        "products",
        ["商品编码", "商品名称", "分类编码", "基础单位编码"],
        [["NEW", "新商品", "CAT", "PCS"]],
    )
    assert new_product.json()["status"] == "PREVIEW_READY"
    changed = await c.put(
        "/api/v1/categories/" + import_refs["CAT"]["id"],
        json={"code": "CAT", "name": "分类修订", "expected_version": 1},
        headers={"Idempotency-Key": uuid4().hex},
    )
    assert changed.status_code == 200
    assert (await confirm(c, new_product.json()["id"])).status_code == 200
    assert (await runner.drain(sessions, identities[0]["org"]))["failed"] == 1
    row = (await preview_rows(c, new_product.json()["id"]))[0]
    assert row["error_code"] == "REFERENCE_CHANGED" and not row["retryable"]


@pytest.mark.parametrize(
    "attributes,expected_code",
    [
        ({"属性:长度": "12.000001", "属性:材质": "steel", "属性:镀锌": "是"}, None),
        ({"属性:长度": "", "属性:材质": "steel", "属性:镀锌": "否"}, "INVALID_ATTRIBUTE"),
        ({"属性:长度": "12", "属性:材质": "unknown", "属性:镀锌": "否"}, "INVALID_ATTRIBUTE"),
        ({"属性:长度": "NaN", "属性:材质": "steel", "属性:镀锌": "否"}, "INVALID_ATTRIBUTE"),
        ({"属性:长度": "12", "属性:材质": "steel", "属性:镀锌": "也许"}, "INVALID_ATTRIBUTE"),
        ({"属性:未知": "不能添加", "属性:长度": "12", "属性:材质": "steel"}, "INVALID_ATTRIBUTE"),
    ],
)
async def test_product_attributes_use_current_catalog_template(
    catalog_client,
    identities,
    import_refs,
    attributes,
    expected_code,
):
    c = catalog_client
    definition = [
        {"key": "长度", "label": "长度", "kind": "decimal", "required": True},
        {
            "key": "材质",
            "label": "材质",
            "kind": "enum",
            "options": ["steel", "copper"],
            "required": True,
        },
        {"key": "镀锌", "label": "镀锌", "kind": "boolean"},
    ]
    category = await create(
        c, "categories", {"code": "ATTR", "name": "带属性的分类", "attribute_schema": definition}
    )
    assert category.status_code == 201
    values = {
        "商品编码": "ATTR_ITEM",
        "商品名称": "属性商品",
        "分类编码": "ATTR",
        "基础单位编码": "PCS",
        **attributes,
    }
    response = await upload(c, "products", list(values), [list(values.values())])
    assert response.status_code == 201, response.text
    rows = await preview_rows(c, response.json()["id"])
    if expected_code:
        assert response.json()["status"] == "PREVIEW_INVALID", rows
        assert any(expected_code == e["code"] for e in rows[0]["errors"]), rows
    else:
        assert response.json()["status"] == "PREVIEW_READY", rows
        assert rows[0]["cleaned_values"]["属性:镀锌"] is True
        assert (await confirm(c, response.json()["id"])).status_code == 200
        assert (await runner.drain(sessions, identities[0]["org"]))["succeeded"] == 1
        final = (await preview_rows(c, response.json()["id"]))[0]
        product = (await c.get("/api/v1/products/" + final["target_id"])).json()
        assert product["attributes"] == {"长度": "12.000001", "材质": "steel", "镀锌": True}


async def test_same_file_new_reference_requires_separate_confirmed_batch(catalog_client):
    result = await upload(
        catalog_client,
        "categories",
        ["编码", "名称", "父分类编码"],
        [["PARENT", "父", ""], ["CHILD", "子", "PARENT"]],
    )
    assert result.status_code == 201
    rows = await preview_rows(catalog_client, result.json()["id"])
    assert rows[0]["status"] == "READY" and rows[1]["error_code"] == "UNKNOWN_REFERENCE"
    assert (await confirm(catalog_client, result.json()["id"])).status_code == 409
    assert (await catalog_client.get("/api/v1/categories")).json()["total"] == 0


@pytest.mark.parametrize(
    "resource,headers,rows",
    [
        (
            "products",
            ["商品编码", "商品名称", "分类编码", "基础单位编码", "条码"],
            [["X", "一", "CAT", "PCS", "001"], ["Y", "二", "CAT", "PCS", "001"]],
        ),
        (
            "supplier-products",
            ["供应商编码", "商品编码", "供应商货号", "采购单位编码"],
            [["SUPPLIER", "ITEM", "SAME", "PCS"], ["SUPPLIER", "ITEM", "SAME", "PCS"]],
        ),
        (
            "product-prices",
            ["商品编码", "价格类型", "价格"],
            [["ITEM", "standard", "10"], ["ITEM", "standard", "20"]],
        ),
    ],
)
async def test_file_duplicate_secondary_and_price_keys_all_reported(
    catalog_client, import_refs, resource, headers, rows
):
    result = await upload(catalog_client, resource, headers, rows)
    assert result.status_code == 201
    records = await preview_rows(catalog_client, result.json()["id"])
    assert len(records) == 2
    assert all("DUPLICATE_IN_FILE" in {e["code"] for e in r["errors"]} for r in records)


async def test_update_blank_numeric_is_distinct_from_explicit_zero_and_base_unit_is_frozen(
    catalog_client, import_refs
):
    c = catalog_client
    for value, valid in [("", False), ("0", True)]:
        result = await upload(
            c, "products", ["商品编码", "最低库存数量"], [["ITEM", value]], mode="UPDATE_EXISTING"
        )
        assert result.status_code == 201
        assert result.json()["status"] == ("PREVIEW_READY" if valid else "PREVIEW_INVALID")
    base = await upload(
        c, "products", ["商品编码", "基础单位编码"], [["ITEM", "BOX"]], mode="UPDATE_EXISTING"
    )
    assert base.status_code == 201
    assert base.json()["status"] == "PREVIEW_INVALID"
    error = (await preview_rows(c, base.json()["id"]))[0]
    assert "BASE_UNIT_IMMUTABLE" in {e["code"] for e in error["errors"]}


async def test_bounded_upload_errors_close_spooled_files_and_never_create_batch(
    catalog_client, monkeypatch
):
    import starlette.formparsers

    c = catalog_client
    opened = []
    original = starlette.formparsers.SpooledTemporaryFile

    def tracked(*args, **kwargs):
        handle = original(*args, **kwargs)
        opened.append(handle)
        return handle

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", tracked)
    invalid = await upload(c, "brands", [], [], data=b"invalid XLSX " * 100000)
    assert invalid.status_code == 422
    assert opened and all(handle.closed for handle in opened)
    too_many = await c.post(
        BASE + "/previews",
        files=[("file", ("one.xlsx", b"x")), ("file", ("two.xlsx", b"y"))],
        data={"resource": "brands"},
        headers={"Idempotency-Key": uuid4().hex},
    )
    assert too_many.status_code == 400
    assert all(handle.closed for handle in opened)
    unknown = await c.post(
        BASE + "/previews",
        files={"file": ("one.xlsx", workbook_bytes(["编码", "名称"], [["X", "X"]]))},
        data={"resource": "brands", "organization_id": str(uuid4())},
        headers={"Idempotency-Key": uuid4().hex},
    )
    assert unknown.status_code == 422 and unknown.json()["code"] == "INVALID_FORM"
    assert all(handle.closed for handle in opened)
    assert (await c.get(BASE)).json()["total"] == 0


async def test_malformed_multipart_is_problem_details_and_closes_file(catalog_client, monkeypatch):
    import starlette.formparsers

    opened = []
    original = starlette.formparsers.SpooledTemporaryFile

    def tracked(*args, **kwargs):
        handle = original(*args, **kwargs)
        opened.append(handle)
        return handle

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", tracked)
    body = (
        b'--broken\r\nContent-Disposition: form-data; name="file"; filename="x.xlsx"\r\n'
        b"Content-Type: application/octet-stream\r\n\r\npayload\r\n"
        b"--broken\r\nInvalid Header Without Colon\r\n\r\nbody\r\n--broken--\r\n"
    )
    response = await catalog_client.post(
        BASE + "/previews",
        content=body,
        headers={
            "Content-Type": "multipart/form-data; boundary=broken",
            "Idempotency-Key": uuid4().hex,
        },
    )
    assert response.status_code == 422 and response.json()["code"] == "INVALID_FORM"
    assert response.headers["content-type"].startswith("application/problem+json")
    assert opened and all(handle.closed for handle in opened)


async def test_failure_during_large_preview_rolls_back_the_entire_unconfirmed_batch(
    catalog_client,
    identities,
    monkeypatch,
):
    from sqlalchemy import text
    from sqlalchemy.ext.asyncio import AsyncSession

    from forge_erp.core.db import set_tenant

    original = AsyncSession.execute
    groups = 0

    async def interrupted(self, statement, *args, **kwargs):
        nonlocal groups
        result = await original(self, statement, *args, **kwargs)
        if "INSERT INTO forge.import_rows" in str(statement):
            groups += 1
            if groups == 2:
                raise RuntimeError("lost connection while persisting a preview")
        return result

    with monkeypatch.context() as patch:
        patch.setattr(AsyncSession, "execute", interrupted)
        response = await upload(
            catalog_client,
            "brands",
            ["编码", "名称"],
            [["B" + str(i), "独立资料"] for i in range(501)],
        )
    assert groups == 2 and response.status_code == 500
    assert (await catalog_client.get(BASE)).json()["total"] == 0
    assert (await catalog_client.get("/api/v1/brands")).json()["total"] == 0
    async with sessions.begin() as db:
        await set_tenant(db, identities[0]["org"])
        assert (await db.execute(text("SELECT count(*) FROM forge.import_rows"))).scalar_one() == 0
        assert (
            await db.execute(
                text(
                    "SELECT count(*) FROM forge.audit_events WHERE action='catalog.import.preview'"
                )
            )
        ).scalar_one() == 0
        assert (
            await db.execute(
                text(
                    "SELECT count(*) FROM forge.outbox_events "
                    "WHERE event_type='catalog.import.preview'"
                )
            )
        ).scalar_one() == 0
