"""Bounded OOXML parsing; numeric values never pass through openpyxl's float decoder."""

import hashlib
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Never

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import get_column_letter

from forge_erp.core.errors import Problem
from forge_erp.modules.catalog_import.domain.columns import COLUMNS

MAX_UPLOAD = 10_000_000
MAX_EXPANDED = 100_000_000
MAX_MEMBER = 50_000_000
MAX_ROWS = 10_000
MAX_COLUMNS = 64
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


@dataclass(frozen=True)
class Cell:
    value: str
    kind: str


@dataclass(frozen=True)
class Sheet:
    filename: str
    file_hash: str
    worksheet: str
    headers: list[str]
    rows: list[tuple[int, dict[str, Cell]]]


def invalid(code, message) -> Never:
    raise Problem(422, code, message)


def local(tag):
    return tag.rsplit("}", 1)[-1]


def xml(data):
    try:
        root = fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        stack = [(root, 1)]
        while stack:
            element, depth = stack.pop()
            if depth > 64:
                raise ValueError("XML nesting exceeds 64")
            stack.extend((child, depth + 1) for child in element)
        return root
    except Exception as exc:
        raise Problem(422, "UNSAFE_XML", "工作簿包含不安全或损坏的 XML") from exc


def archive(data):
    if len(data) > MAX_UPLOAD:
        raise Problem(413, "FILE_TOO_LARGE", "文件最多 10,000,000 字节")
    parts = {}
    expanded = 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            if len(package.infolist()) > 1024:
                invalid("ZIP_LIMIT", "工作簿 ZIP 成员过多")
            for info in package.infolist():
                name = info.filename
                if (
                    "\\" in name
                    or name.startswith("/")
                    or ":" in name
                    or ".." in PurePosixPath(name).parts
                    or name in parts
                    or posixpath.normpath(name).rstrip("/") != name.rstrip("/")
                ):
                    invalid("UNSAFE_ZIP_PATH", "工作簿包含危险或重复文件路径")
                if info.is_dir():
                    parts[name] = b""
                    continue
                if info.flag_bits & 1:
                    invalid("ENCRYPTED_WORKBOOK", "不接受加密工作簿")
                lower = name.lower()
                if any(
                    x in lower
                    for x in (
                        "vbaproject",
                        "macrosheet",
                        "externallink",
                        "embeddings/",
                        "activex/",
                        "connections.xml",
                        "querytables/",
                    )
                ):
                    invalid("UNSAFE_WORKBOOK", "不接受宏、嵌入对象或外部数据链接")
                if info.file_size > MAX_MEMBER or info.file_size > max(info.compress_size, 1) * 100:
                    invalid("ZIP_LIMIT", "工作簿成员大小或压缩比超过限制")
                chunks = []
                actual = 0
                with package.open(info) as member:
                    while block := member.read(65536):
                        actual += len(block)
                        expanded += len(block)
                        if (
                            actual > MAX_MEMBER
                            or expanded > MAX_EXPANDED
                            or actual > max(info.compress_size, 1) * 100
                        ):
                            invalid("ZIP_LIMIT", "实际解压字节超过工作簿限制")
                        chunks.append(block)
                content = b"".join(chunks)
                if lower.endswith((".xml", ".rels")):
                    root = xml(content)
                    for item in root.iter():
                        tag = local(item.tag)
                        if tag in {"f", "definedName"}:
                            invalid(
                                "FORMULA_NOT_ALLOWED", "不接受公式或公式定义；请粘贴为值后再导入"
                            )
                        if (
                            tag == "Relationship"
                            and item.get("TargetMode", "").lower() == "external"
                        ):
                            invalid("EXTERNAL_LINK_NOT_ALLOWED", "不接受外部链接或超链接")
                        if tag in {"Override", "Default"} and any(
                            x in item.get("ContentType", "").lower()
                            for x in ("macroenabled", "vbaproject", "oleobject")
                        ):
                            invalid("UNSAFE_WORKBOOK", "不接受宏或嵌入对象工作簿")
                parts[name] = content
    except Problem:
        raise
    except (
        zipfile.BadZipFile,
        RuntimeError,
        OSError,
        EOFError,
        NotImplementedError,
        ValueError,
    ) as exc:
        raise Problem(422, "INVALID_XLSX", "不是可读取的未加密 xlsx 文件") from exc
    if "xl/workbook.xml" not in parts or "[Content_Types].xml" not in parts:
        invalid("INVALID_XLSX", "缺少工作簿结构")
    return parts


def text_content(element):
    if element is None:
        return ""
    result, stack = [], [element]
    while stack:
        item = stack.pop()
        if local(item.tag) == "rPh":
            continue
        if local(item.tag) == "t":
            result.append(item.text or "")
        else:
            stack.extend(reversed(list(item)))
    return "".join(result)


def parse(data: bytes, filename: str, worksheet: str | None = None) -> Sheet:
    if not filename.lower().endswith(".xlsx"):
        invalid("INVALID_XLSX", "仅接受 .xlsx 文件")
    parts = archive(data)
    workbook = xml(parts["xl/workbook.xml"])
    relations = xml(parts.get("xl/_rels/workbook.xml.rels", b"<Relationships/>"))
    targets = {r.get("Id"): r.get("Target", "") for r in relations}
    sheets = list(workbook.iter(NS + "sheet"))
    names = [item.get("name", "") for item in sheets]
    if not worksheet:
        if len(names) != 1:
            invalid("WORKSHEET_REQUIRED", "请选择一个工作表：" + "、".join(names)[:1000])
        worksheet = names[0]
    if names.count(worksheet) != 1:
        invalid("UNKNOWN_WORKSHEET", "指定工作表不存在或名称不唯一")
    target = targets.get(sheets[names.index(worksheet)].get(REL + "id"), "")
    sheet_path = (
        target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    )
    if sheet_path not in parts or not sheet_path.startswith("xl/worksheets/"):
        invalid("INVALID_XLSX", "工作表关系不合法")
    shared = []
    if "xl/sharedStrings.xml" in parts:
        shared = [text_content(si) for si in xml(parts["xl/sharedStrings.xml"])]
    date_styles = set()
    if "xl/styles.xml" in parts:
        style = xml(parts["xl/styles.xml"])
        formats = dict(BUILTIN_FORMATS)
        for item in style.iter(NS + "numFmt"):
            formats[int(item.get("numFmtId", "0"))] = item.get("formatCode", "")
        cell_styles = style.find(NS + "cellXfs")
        for index, item in enumerate([] if cell_styles is None else cell_styles):
            if is_date_format(formats.get(int(item.get("numFmtId", "0")), "")):
                date_styles.add(str(index))
    rows = []
    logical_size = 0
    seen_rows = set()
    for row in xml(parts[sheet_path]).iter(NS + "row"):
        try:
            row_no = int(row.get("r", "0"))
        except ValueError:
            invalid("INVALID_XLSX", "工作表行号不合法")
        if row_no < 1 or row_no in seen_rows:
            invalid("INVALID_XLSX", "工作表行号重复或不合法")
        seen_rows.add(row_no)
        cells = {}
        for item in row.findall(NS + "c"):
            match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", item.get("r", ""))
            if match is None or int(match[2]) != row_no:
                invalid("INVALID_XLSX", "单元格坐标不合法")
            column = 0
            for char in match[1]:
                column = column * 26 + ord(char) - ord("A") + 1
            if column > MAX_COLUMNS or column in cells:
                invalid("COLUMN_LIMIT", "最多 64 列，单元格坐标不能重复")
            kind = item.get("t", "n")
            value = item.findtext(NS + "v", "")
            if kind == "s":
                try:
                    index = int(value)
                    if index < 0:
                        raise IndexError
                    value = shared[index]
                except ValueError, IndexError:
                    invalid("INVALID_XLSX", "共享字符串索引不合法")
                kind = "text"
            elif kind == "inlineStr":
                value = text_content(item.find(NS + "is"))
                kind = "text"
            elif kind == "str":
                kind = "text"
            elif kind == "b":
                kind = "boolean"
            elif kind == "d" or (kind == "n" and item.get("s", "0") in date_styles):
                kind = "date"
            elif kind == "n":
                kind = "number"
            else:
                kind = "error"
            logical_size += len(value.encode("utf-8"))
            if logical_size > MAX_EXPANDED or len(value) > 32767:
                invalid("CELL_CONTENT_LIMIT", "单元格正文过长或展开后总量超过 100MB")
            cells[column] = Cell(value, kind)
        if any(cell.value != "" for cell in cells.values()):
            rows.append((row_no, cells))
            if len(rows) > MAX_ROWS + 1:
                invalid("ROW_LIMIT", "最多 10,000 条非空数据行")
    if not rows or rows[0][0] != 1:
        invalid("HEADERS_REQUIRED", "第一行必须为模板列名")
    header_cells = rows.pop(0)[1]
    headers = [
        header_cells.get(index, Cell("", "text")).value.strip()
        for index in range(1, max(header_cells) + 1)
    ]
    if any(not value for value in headers) or len(set(headers)) != len(headers):
        invalid("INVALID_HEADERS", "不接受空白或重复列名")
    result = []
    for row_no, cells in rows:
        if any(index > len(headers) and cell.value for index, cell in cells.items()):
            invalid("UNKNOWN_COLUMN", "数据出现在没有列名的位置")
        result.append(
            (
                row_no,
                {
                    header: cells.get(index, Cell("", "text"))
                    for index, header in enumerate(headers, 1)
                },
            )
        )
    if not result:
        invalid("EMPTY_WORKSHEET", "工作表没有需要导入的数据行")
    return Sheet(
        filename.replace("\\", "/").rsplit("/", 1)[-1][:200],
        hashlib.sha256(data).hexdigest(),
        worksheet,
        headers,
        result,
    )


def workbook_bytes(headers, rows, comments=None):
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "资料"
    for values in [headers, *rows]:
        sheet.append(["" if value is None else str(value) for value in values])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
    if comments:
        for index, note in enumerate(comments, 1):
            sheet.cell(1, index).comment = Comment(note, "Forge ERP")
    sheet.freeze_panes = "A2"
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def template(resource):
    from forge_erp.modules.catalog_import.domain.columns import KEYS

    notes = []
    for item in COLUMNS[resource].values():
        note = "新增必填。" if item.required else "新增可选；空值按现有资料规则处理。"
        if item.field in KEYS[resource]:
            note += "用于定位资料的自然键；更新必须提供。"
        if item.reference:
            note += "引用当前组织已启用资料编码：" + item.reference + "。"
        if item.kind == "identifier":
            note += "必须使用文本单元格，保留前导零与大小写；禁止数字编码。"
        if item.kind == "decimal":
            note += "精确十进制，最多6位小数；不支持公式、布尔或日期。"
        if item.kind == "integer":
            note += "0至3650的整数天数。"
        if item.kind == "json":
            note += "分类属性模板JSON数组，沿用分类编辑的字段规则。"
        if item.field == "price_type":
            note += "standard / retail / wholesale / customer；customer须提供客户编码。"
        if item.field == "price_tier":
            note += "standard / retail / wholesale；留空默认standard。"
        note += "更新时未映射列保留原值；映射空值不会自动保留，请检查完整清洗值。"
        notes.append(note)
    return workbook_bytes(list(COLUMNS[resource]), [], notes)
